from docx import *

doc = Document('2745.docx')
tables = doc.tables
table = tables[1]
celda = table.cell(19,1)
celda.text = "este es un test de como se hace jajajajaja"
doc.save('hola.docx')

